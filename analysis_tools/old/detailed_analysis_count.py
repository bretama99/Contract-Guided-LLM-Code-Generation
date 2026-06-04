"""
Currently counts error types and summarizes counts
"""


import argparse, re
from pathlib import Path
from openpyxl import Workbook, load_workbook
from collections import Counter
from analysis_tools.old.analyze_results import HELPED, REGRESSED, FAILED
from formatting import format_worksheet
from src.common.config import DATASETS
from src.common.llm_clients import PROVIDERS, default_model


def summarize(root, dataset, model):
    # Get analysis file
    analysis = get_input_path(root, dataset, model)

    if not analysis.exists():
        raise FileNotFoundError(f"{dataset}: missing analysis file: {analysis}")

    # Open source workbook
    source_wb = load_workbook(analysis, data_only=False)

    # Make new workbook
    target_wb = Workbook()
    target_wb.remove(target_wb.active)

    # Create summary sheet first
    summary_counts = {}
    summary_counts[HELPED + ": Failed Vanilla Only"] = count_failure_errors(source_wb[HELPED], "Vanilla")
    summary_counts[REGRESSED + ": Failed Contract-Guided Only"] = count_failure_errors(source_wb[REGRESSED], "Contract-Guided")
    summary_counts[FAILED + ": Vanilla"] = count_failure_errors(source_wb[FAILED], "Vanilla")
    summary_counts[FAILED + ": Contract-Guided"] = count_failure_errors(source_wb[FAILED], "Contract-Guided")

    write_summary_sheet(target_wb, summary_counts)

    # Copy other sheets into the same target workbook
    for sheet_name in [HELPED, REGRESSED, FAILED]:
        copy_sheet_values(source_wb, target_wb, sheet_name)

    # Output path
    if dataset == "evalplus":
        fname = "EvalPlus_detailed_analysis_counts.xlsx"
    elif dataset == "evalplus_mbpp":
        fname = "EvalPlus_MBPP_detailed_analysis_counts.xlsx"
    else:
        fname = f"{dataset}_detailed_analysis_counts.xlsx"

    out = root / "analysis" / safe(model) / "detailed" / fname
    out.parent.mkdir(parents=True, exist_ok=True)

    # Save new workbook
    target_wb.save(out)
    print(f"Wrote {out}")


def write_summary_sheet(target_wb, summary_counts, summary_sheet_name="Failure Error Summary"):
    
    summary_ws = target_wb.create_sheet(summary_sheet_name)

    summary_ws.append(["Sheet Name", "Failure Error", "Count"])

    for sheet_name, counts in summary_counts.items():
        for error, count in counts.items():
            summary_ws.append([sheet_name, error, count])
        summary_ws.append([])

    format_worksheet(summary_ws)


def count_failure_errors(worksheet, failure_type):
    
    error_col = find_column_by_header(worksheet, f"{failure_type} Error Type")

    if error_col is None:
        raise ValueError(
            f"No column named '{failure_type} Error Type' in sheet '{worksheet.title}'"
        )

    sheet_counts = Counter()

    for row in range(2, worksheet.max_row + 1):
        value = worksheet.cell(row=row, column=error_col).value

        if value is None:
            continue

        error_text = str(value).strip()

        if error_text == "":
            continue

        error_categories = error_text.split(",")

        for category in error_categories:
            category = category.strip()

            if category == "":
                continue

            sheet_counts[category] += 1

    return sheet_counts


def copy_sheet_values(source_wb, target_wb, sheet_name):
    source_ws = source_wb[sheet_name]
    target_ws = target_wb.create_sheet(sheet_name)

    for row in source_ws.iter_rows(values_only=True):
        target_ws.append(row)

    format_worksheet(target_ws)


def safe(x):
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", str(x)).strip("._-") or "x"


def find_column_by_header(ws, header_name):
    for cell in ws[1]:
        if cell.value and str(cell.value).strip().lower() == header_name.lower():
            return cell.column
    return None


def get_input_path(root, dataset, model):

    if dataset == "evalplus":
        fname = "EvalPlus_detailed_analysis.xlsx"
    elif dataset == "evalplus_mbpp":
        fname = "EvalPlus_MBPP_detailed_analysis.xlsx"
    else:
        fname = f"{dataset}_detailed_analysis.xlsx"

    return root / "analysis" / safe(model) / "detailed" / fname


def main():

    p = argparse.ArgumentParser()
    p.add_argument("--root", default=".")
    p.add_argument("--dataset", choices=sorted(DATASETS), required=True)
    p.add_argument("--provider", choices=sorted(PROVIDERS), required=True)
    p.add_argument("--model")
    args = p.parse_args()

    root = Path(args.root).resolve()
    dataset = args.dataset
    model = args.model or default_model(args.provider)

    summarize(root, dataset, model)


if __name__ == "__main__":
    main()