from openpyxl.styles import Font, Alignment, Side, Border


def format_worksheet(worksheet) -> None:
    worksheet.sheet_view.showGridLines = True
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.zoomScale = 75

    header_font = Font(name="Times New Roman", size=14, bold=True)
    body_font = Font(name="Aptos", size=12)

    thin_gray = Side(style="thin", color="D9D9D9")
    border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray,
    )

    max_row = worksheet.max_row
    max_col = worksheet.max_column

    for row in worksheet.iter_rows(
        min_row=1,
        max_row=max_row,
        min_col=1,
        max_col=max_col,
    ):
        for cell in row:
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(
                horizontal=_default_horizontal_alignment(cell.value),
                vertical="top",
                wrap_text=True,
            )

    _format_header_row(
        worksheet,
        row_number=1,
        header_font=header_font,
        border=border,
    )

    _size_columns(worksheet)


def format_summary_worksheet(worksheet) -> None:
    worksheet.sheet_view.showGridLines = True
    worksheet.freeze_panes = "A1"
    worksheet.sheet_view.zoomScale = 75

    header_font = Font(name="Times New Roman", size=14, bold=True)
    body_font = Font(name="Aptos", size=12)
    index_font = Font(name="Times New Roman", size=14, bold=True)

    thin_gray = Side(style="thin", color="D9D9D9")
    border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray,
    )

    max_row = worksheet.max_row
    max_col = worksheet.max_column

    for row in worksheet.iter_rows(
        min_row=1,
        max_row=max_row,
        min_col=1,
        max_col=max_col,
    ):
        for cell in row:
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(
                horizontal=_default_horizontal_alignment(cell.value),
                vertical="top",
                wrap_text=True,
            )

    # Overall summary:
    # written with startrow=0, startcol=0, index=True
    # header row is Excel row 1, columns A:C.
    _format_header_range(
        worksheet,
        row_number=1,
        min_col=1,
        max_col=3,
        header_font=header_font,
        border=border,
    )

    # Intersections summary:
    # written with startrow=0, startcol=4, index=True
    # header row is Excel row 1, columns E:J.
    _format_header_range(
        worksheet,
        row_number=1,
        min_col=5,
        max_col=10,
        header_font=header_font,
        border=border,
    )

    # Error buckets summary:
    # written with startrow=5, startcol=4, index=True
    # header row is Excel row 6, columns E:J.
    _format_header_range(
        worksheet,
        row_number=6,
        min_col=5,
        max_col=10,
        header_font=header_font,
        border=border,
    )

    # Bold index labels for the overall summary table.
    _format_index_range(
        worksheet,
        min_row=2,
        max_row=8,
        column=1,
        index_font=index_font,
        border=border,
    )

    # Bold index labels for the contract component summary table.
    _format_index_range(
        worksheet,
        min_row=10,
        max_row=12,
        column=1,
        index_font=index_font,
        border=border,
    )

    # Bold index labels for the effect summary table, if present.
    if worksheet.max_row >= 14:
        _format_index_range(
            worksheet,
            min_row=14,
            max_row=worksheet.max_row,
            column=1,
            index_font=index_font,
            border=border,
        )

    # Bold index labels for the correctness/completeness intersection table.
    _format_index_range(
        worksheet,
        min_row=2,
        max_row=3,
        column=5,
        index_font=index_font,
        border=border,
    )

    # Bold index labels for the error bucket summary table.
    _format_index_range(
        worksheet,
        min_row=7,
        max_row=9,
        column=5,
        index_font=index_font,
        border=border,
    )

    _size_columns(worksheet)

    # Make the summary row-label column wider.
    worksheet.column_dimensions["A"].width = 40

    # Show percentage-like values as whole-number percentages.
    # Values are already stored as 0-100 numbers, so use 0"%" instead of 0%.
    _format_summary_percentages(worksheet)


def format_analysis_summary_worksheet(worksheet) -> None:
    """Format the vertically stacked Summary sheet used by contract analysis.

    This function is intentionally separate from format_summary_worksheet(),
    because the older formatter is tied to a different fixed-position summary
    layout. The contract-analysis Summary sheet writes several pandas tables
    vertically, with blank rows between tables, so this formatter detects each
    table header dynamically and bolds all column names.
    """
    worksheet.sheet_view.showGridLines = True
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.zoomScale = 75

    header_font = Font(name="Times New Roman", size=14, bold=True)
    body_font = Font(name="Aptos", size=12)

    thin_gray = Side(style="thin", color="D9D9D9")
    border = Border(
        left=thin_gray,
        right=thin_gray,
        top=thin_gray,
        bottom=thin_gray,
    )

    max_row = worksheet.max_row
    max_col = worksheet.max_column

    for row in worksheet.iter_rows(
        min_row=1,
        max_row=max_row,
        min_col=1,
        max_col=max_col,
    ):
        for cell in row:
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(
                horizontal=_default_horizontal_alignment(cell.value),
                vertical="top",
                wrap_text=True,
            )

    for row_number in _analysis_summary_header_rows(worksheet):
        last_col = _last_non_empty_column(worksheet, row_number)
        if last_col is None:
            continue

        _format_summary_header_range(
            worksheet,
            row_number=row_number,
            min_col=1,
            max_col=last_col,
            header_font=header_font,
            border=border,
        )

    _size_columns(worksheet)
    _format_analysis_summary_percentages(worksheet)


def _format_header_row(
    worksheet,
    row_number: int,
    header_font: Font,
    border: Border,
) -> None:
    if row_number < 1 or row_number > worksheet.max_row:
        return

    for cell in worksheet[row_number]:
        cell.font = header_font
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


def _format_header_range(
    worksheet,
    row_number: int,
    min_col: int,
    max_col: int,
    header_font: Font,
    border: Border,
) -> None:
    if row_number < 1 or row_number > worksheet.max_row:
        return

    for row in worksheet.iter_rows(
        min_row=row_number,
        max_row=row_number,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )


def _format_summary_header_range(
    worksheet,
    row_number: int,
    min_col: int,
    max_col: int,
    header_font: Font,
    border: Border,
) -> None:
    if row_number < 1 or row_number > worksheet.max_row:
        return

    for row in worksheet.iter_rows(
        min_row=row_number,
        max_row=row_number,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=True,
            )


def _format_index_range(
    worksheet,
    min_row: int,
    max_row: int,
    column: int,
    index_font: Font,
    border: Border,
) -> None:
    if min_row < 1 or min_row > worksheet.max_row:
        return

    max_row = min(max_row, worksheet.max_row)

    for row in worksheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=column,
        max_col=column,
    ):
        for cell in row:
            cell.font = index_font
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left",
                vertical="top",
                wrap_text=True,
            )


def _analysis_summary_header_rows(worksheet) -> list[int]:
    """Return header rows for vertically stacked pandas tables.

    A header row is the first non-empty row after the worksheet start or after
    one or more completely blank rows.
    """
    header_rows = []
    previous_row_was_blank = True

    for row_number in range(1, worksheet.max_row + 1):
        row_has_value = any(
            worksheet.cell(row=row_number, column=col).value not in (None, "")
            for col in range(1, worksheet.max_column + 1)
        )

        if row_has_value and previous_row_was_blank:
            header_rows.append(row_number)

        previous_row_was_blank = not row_has_value

    return header_rows


def _last_non_empty_column(worksheet, row_number: int) -> int | None:
    for col in range(worksheet.max_column, 0, -1):
        if worksheet.cell(row=row_number, column=col).value not in (None, ""):
            return col

    return None


def _format_summary_percentages(worksheet) -> None:
    for row in worksheet.iter_rows():
        for cell in row:
            if not isinstance(cell.value, (int, float)) or isinstance(cell.value, bool):
                continue

            header = None

            # Left-side summary tables use column C for percentages.
            if cell.column == 3:
                header = worksheet.cell(row=1, column=3).value

            # Error buckets table uses columns G:J for percentage columns.
            elif cell.row >= 7 and cell.column in (7, 8, 9, 10):
                header = worksheet.cell(row=6, column=cell.column).value

            if isinstance(header, str) and (
                "percentage" in header.lower()
                or header.strip().startswith("%")
            ):
                cell.number_format = '0"%"'


def _format_analysis_summary_percentages(worksheet) -> None:
    """Format percentage columns in dynamically positioned summary tables.

    Values in these tables are stored as 0-100 numbers, so the Excel number
    format uses 0"%" instead of 0%.
    """
    header_rows = set(_analysis_summary_header_rows(worksheet))
    current_headers = {}

    for row_number in range(1, worksheet.max_row + 1):
        if row_number in header_rows:
            current_headers = {
                col: worksheet.cell(row=row_number, column=col).value
                for col in range(1, worksheet.max_column + 1)
            }
            continue

        for col in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(row=row_number, column=col)

            if not isinstance(cell.value, (int, float)) or isinstance(cell.value, bool):
                continue

            header = current_headers.get(col)

            if isinstance(header, str) and (
                "percentage" in header.lower()
                or "percent" in header.lower()
                or header.strip().startswith("%")
            ):
                cell.number_format = '0"%"'


def _size_columns(worksheet) -> None:
    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter

        for cell in column_cells:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        adjusted_width = min(max(max_length * 1.2, 15), 70)
        worksheet.column_dimensions[column_letter].width = adjusted_width


def _default_horizontal_alignment(value) -> str:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "right"

    return "left"