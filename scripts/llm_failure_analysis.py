from __future__ import annotations

import argparse
import ast
import json
import re
import time
from collections import Counter
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from src.common.llm_clients import call_chat_model, get_client

JsonDict = dict[str, Any]

BENCHMARK = "classeval"

SOURCE_RUNS = [
    {
        "label": "OpenRouter GPT-3.5",
        "provider": "openrouter",
        "model": "openai-gpt-3.5-turbo",
        "model_aliases": [
            "openai-gpt-3.5-turbo",
            "openai_gpt-3.5-turbo",
            "openai/gpt-3.5-turbo",
        ],
    },
    {
        "label": "Qwen2.5-32B",
        "provider": "qwen",
        "model": "qwen2.5-32b-instruct",
        "model_aliases": [
            "qwen2.5-32b-instruct",
            "qwen2.5_32b_instruct",
        ],
    },
]

ANALYZER_PROVIDER = "openrouter"
ANALYZER_MODEL = "openai/gpt-5.5"

PIPELINES = [
    ("Vanilla", "vanilla"),
    ("Contract-guided", "contract_guided"),
    ("Feedback-optimized", "rl_optimized_contract_guided"),
]

ERROR_TYPES = [
    "Logical Error",
    "Runtime Error",
    "Syntax Error",
    "Generation Error",
    "Timeout Error",
    "Other Error",
]

CONTRACT_TYPES = [
    "interface",
    "precondition",
    "postcondition",
    "invariant",
    "edge_case",
    "invalid_input_behavior",
    "dependency_interaction",
    "not_contract_related",
]

DETAIL_COLUMNS = [
    "Source Model",
    "Pipeline",
    "Task ID",
    "Class Name",
    "Methods List",
    "Passed",
    "Tests Passed",
    "Tests Total",
    "Failed Method / Test",
    "Expected Output",
    "Actual Output",
    "Error Type",
    "Raw Failure Type",
    "Contract Type",
    "Reason for Failure",
    "How to Solve",
]


def text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def as_dict(value: Any) -> JsonDict:
    return value if isinstance(value, dict) else {}


def as_list(value: Any) -> list[Any]:
    return value if isinstance(value, list) else []


def safe_name(value: Any) -> str:
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", str(value)).strip("._-") or "unnamed"


def short(value: Any, limit: int = 900) -> str:
    if value is None:
        return ""
    if isinstance(value, BaseException):
        value = str(value)
    elif not isinstance(value, str):
        try:
            value = json.dumps(value, ensure_ascii=False)
        except TypeError:
            value = str(value)
    value = value.strip()
    return value if len(value) <= limit else value[:limit] + "\n...[truncated]"


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def resolve(root: Path, raw: Any) -> Path | None:
    value = text(raw)
    if not value:
        return None

    path = Path(value)
    if path.exists():
        return path

    for anchor in ("outputs", "results", "data"):
        if anchor in path.parts:
            candidate = root.joinpath(*path.parts[path.parts.index(anchor):])
            if candidate.exists():
                return candidate

    candidate = root / value
    return candidate if candidate.exists() else None


def find_dataset(root: Path, explicit: str | None) -> Path:
    if explicit:
        path = resolve(root, explicit)
        if path:
            return path
        raise FileNotFoundError(explicit)

    candidates = [
        root / "data" / "processed" / BENCHMARK / "ClassEval_data.json",
        root / "data" / "processed" / BENCHMARK / f"{BENCHMARK}_data.json",
        root / "data" / "processed" / BENCHMARK / f"{BENCHMARK}.json",
    ]
    for path in candidates:
        if path.exists():
            return path

    raise FileNotFoundError("Dataset file not found. Use --dataset-file.")


def task_id_of(task: JsonDict) -> str:
    for key in ("task_id", "id", "name", "entry_point"):
        if text(task.get(key)):
            return text(task.get(key))
    return ""


def load_dataset(root: Path, explicit: str | None) -> dict[str, JsonDict]:
    data = load_json(find_dataset(root, explicit))
    tasks = data if isinstance(data, list) else data.get("tasks") or data.get("data") or list(data.values())
    return {task_id_of(t): t for t in tasks if isinstance(t, dict) and task_id_of(t)}


def details_candidates(root: Path, folder: str) -> list[Path]:
    bases = [
        root / "results" / BENCHMARK / folder,
        root / "outputs" / "results" / BENCHMARK / folder,
    ]

    paths: list[Path] = []
    for base in bases:
        if base.exists():
            paths.extend(base.rglob("*details.json"))
    return paths


def matches_model(path: Path, provider: str, aliases: list[str]) -> bool:
    name = path.name
    if safe_name(provider) not in name:
        return False
    return any(safe_name(alias) in name for alias in aliases)


def find_details(root: Path, folder: str, run: JsonDict, required: bool = True) -> Path | None:
    provider = text(run["provider"])
    aliases = [text(x) for x in as_list(run.get("model_aliases"))] or [text(run["model"])]

    hits = [
        path
        for path in details_candidates(root, folder)
        if matches_model(path, provider, aliases)
    ]

    if hits:
        return sorted(hits, key=lambda p: (p.stat().st_mtime, str(p)))[-1]

    if required:
        raise FileNotFoundError(f"Missing {folder} details for {provider}/{aliases}")

    return None


def methods(task: JsonDict) -> list[str]:
    names = [
        text(m.get("method_name"))
        for m in as_list(task.get("methods_info"))
        if isinstance(m, dict) and text(m.get("method_name"))
    ]
    if names:
        return names

    source = text(task.get("skeleton") or task.get("class_constructor"))
    try:
        tree = ast.parse(source)
        return [node.name for node in ast.walk(tree) if isinstance(node, ast.FunctionDef)]
    except SyntaxError:
        return re.findall(r"def\s+([A-Za-z_][A-Za-z0-9_]*)\s*\(", source)


def metrics_of(row: JsonDict) -> JsonDict:
    return as_dict(as_dict(row.get("evaluation")).get("metrics"))


def dataset_test_count(task: JsonDict) -> int:
    code = text(task.get("test"))
    if not code:
        return 0

    try:
        tree = ast.parse(code)
        return sum(1 for node in ast.walk(tree) if isinstance(node, ast.FunctionDef) and node.name.startswith("test"))
    except SyntaxError:
        return len(re.findall(r"def\s+test[A-Za-z0-9_]*\s*\(", code))


def pass_bool(row: JsonDict | None) -> bool:
    return bool(row and row.get("passed") is True)


def by_task(rows: list[JsonDict]) -> dict[str, JsonDict]:
    return {text(row.get("task_id")): row for row in rows if isinstance(row, dict) and text(row.get("task_id"))}


def pipeline_basic(rows: list[JsonDict]) -> JsonDict:
    total = len(rows)
    passed = sum(row.get("passed") is True for row in rows)
    tests_total = 0
    tests_passed = 0

    for row in rows:
        metrics = metrics_of(row)
        tests_total += int(metrics.get("total") or 0)
        tests_passed += int(metrics.get("passed") or 0)

    return {
        "Total Tasks": total,
        "Passed Tasks": passed,
        "Failed Tasks": total - passed,
        "Pass %": round(100 * passed / total, 2) if total else 0,
        "Executed Tests": tests_total,
        "Passed Tests": tests_passed,
        "Test Pass %": round(100 * tests_passed / tests_total, 2) if tests_total else 0,
    }


def raw_to_error_type(raw_type: Any, error_line: str = "") -> str:
    raw = text(raw_type)
    err = text(error_line).lower()

    if raw == "wrong_answer":
        return "Logical Error"

    if raw in {
        "type_error",
        "value_error",
        "key_error",
        "attribute_error",
        "index_error",
        "zero_division_error",
    }:
        return "Runtime Error"

    if raw in {"syntax_error", "import_error", "name_error"}:
        return "Syntax Error"

    if raw in {"generation_failed", "missing_generation"}:
        return "Generation Error"

    if raw == "timeout":
        return "Timeout Error"

    if any(x in err for x in ("syntaxerror", "importerror", "nameerror")):
        return "Syntax Error"

    if any(x in err for x in ("typeerror", "valueerror", "keyerror", "attributeerror", "indexerror", "zerodivisionerror")):
        return "Runtime Error"

    return "Other Error"


def count_detail_field(detail_rows: list[JsonDict], model_label: str, pipeline_label: str, field: str) -> str:
    counts = Counter(
        text(row.get(field)) or "unknown"
        for row in detail_rows
        if row.get("Source Model") == model_label and row.get("Pipeline") == pipeline_label
    )
    return "; ".join(f"{k}={v}" for k, v in sorted(counts.items()))


def main_detail_field(detail_rows: list[JsonDict], model_label: str, pipeline_label: str, field: str) -> str:
    counts = Counter(
        text(row.get(field)) or "unknown"
        for row in detail_rows
        if row.get("Source Model") == model_label and row.get("Pipeline") == pipeline_label
    )
    return counts.most_common(1)[0][0] if counts else ""


def pipeline_summary_rows(
    all_runs: dict[str, dict[str, tuple[str, Path | None, list[JsonDict]]]],
    detail_rows: list[JsonDict],
) -> list[JsonDict]:
    rows: list[JsonDict] = []

    for model_label, loaded in all_runs.items():
        for pipeline_label, (_folder, path, details) in loaded.items():
            rows.append({
                "Section": "Pipeline Performance",
                "Source Model": model_label,
                "Pipeline": pipeline_label,
                **pipeline_basic(details),
                "Main Error Type": main_detail_field(detail_rows, model_label, pipeline_label, "Error Type"),
                "Error Type Counts": count_detail_field(detail_rows, model_label, pipeline_label, "Error Type"),
                "Contract Type Counts": count_detail_field(detail_rows, model_label, pipeline_label, "Contract Type"),
                "Details File": str(path or ""),
            })

    return rows


def comparison_rows(
    all_runs: dict[str, dict[str, tuple[str, Path | None, list[JsonDict]]]]
) -> list[JsonDict]:
    pairs = [
        ("Vanilla", "Contract-guided"),
        ("Vanilla", "Feedback-optimized"),
        ("Contract-guided", "Feedback-optimized"),
    ]
    rows: list[JsonDict] = []

    for model_label, loaded in all_runs.items():
        pipelines = {name: rows_ for name, (_folder, _path, rows_) in loaded.items()}

        for a_name, b_name in pairs:
            a = by_task(pipelines.get(a_name, []))
            b = by_task(pipelines.get(b_name, []))
            tids = sorted(set(a) | set(b))

            a_passed = b_passed = b_improved = b_regressed = both_passed = both_failed = 0

            for tid in tids:
                ap = pass_bool(a.get(tid))
                bp = pass_bool(b.get(tid))

                a_passed += ap
                b_passed += bp
                b_improved += (not ap and bp)
                b_regressed += (ap and not bp)
                both_passed += (ap and bp)
                both_failed += (not ap and not bp)

            rows.append({
                "Section": "Pairwise Comparisons",
                "Source Model": model_label,
                "Comparison": f"{a_name} → {b_name}",
                "Total Comparable Tasks": len(tids),
                "A Passed": a_passed,
                "B Passed": b_passed,
                "B Passed while A Failed": b_improved,
                "A Passed while B Failed": b_regressed,
                "Both Passed": both_passed,
                "Both Failed": both_failed,
                "B Pass Gain": b_passed - a_passed,
            })

    return rows


def three_way_rows(
    all_runs: dict[str, dict[str, tuple[str, Path | None, list[JsonDict]]]]
) -> list[JsonDict]:
    rows: list[JsonDict] = []

    for model_label, loaded in all_runs.items():
        pipelines = {name: rows_ for name, (_folder, _path, rows_) in loaded.items()}
        vanilla = by_task(pipelines.get("Vanilla", []))
        contract = by_task(pipelines.get("Contract-guided", []))
        optimized = by_task(pipelines.get("Feedback-optimized", []))
        tids = sorted(set(vanilla) | set(contract) | set(optimized))

        counts = Counter()

        for tid in tids:
            v = pass_bool(vanilla.get(tid))
            c = pass_bool(contract.get(tid))
            o = pass_bool(optimized.get(tid))

            counts["vanilla_passed"] += v
            counts["contract_passed"] += c
            counts["optimized_passed"] += o
            counts["contract_passed_vanilla_failed"] += c and not v
            counts["vanilla_passed_contract_failed"] += v and not c
            counts["optimized_passed_vanilla_failed"] += o and not v
            counts["vanilla_passed_optimized_failed"] += v and not o
            counts["optimized_passed_contract_failed"] += o and not c
            counts["contract_passed_optimized_failed"] += c and not o
            counts["optimized_passed_both_failed"] += o and not v and not c
            counts["optimized_failed_both_passed"] += not o and v and c
            counts["all_passed"] += v and c and o
            counts["all_failed"] += not v and not c and not o
            counts["only_vanilla"] += v and not c and not o
            counts["only_contract"] += c and not v and not o
            counts["only_optimized"] += o and not v and not c

        rows.append({
            "Section": "Three-way Gains",
            "Source Model": model_label,
            "Total Tasks": len(tids),
            "Vanilla Passed": counts["vanilla_passed"],
            "Contract Passed": counts["contract_passed"],
            "Optimized Passed": counts["optimized_passed"],
            "Contract Passed while Vanilla Failed": counts["contract_passed_vanilla_failed"],
            "Vanilla Passed while Contract Failed": counts["vanilla_passed_contract_failed"],
            "Optimized Passed while Vanilla Failed": counts["optimized_passed_vanilla_failed"],
            "Vanilla Passed while Optimized Failed": counts["vanilla_passed_optimized_failed"],
            "Optimized Passed while Contract Failed": counts["optimized_passed_contract_failed"],
            "Contract Passed while Optimized Failed": counts["contract_passed_optimized_failed"],
            "Optimized Passed while Both Vanilla and Contract Failed": counts["optimized_passed_both_failed"],
            "Optimized Failed while Both Vanilla and Contract Passed": counts["optimized_failed_both_passed"],
            "All Passed": counts["all_passed"],
            "All Failed": counts["all_failed"],
            "Only Vanilla Passed": counts["only_vanilla"],
            "Only Contract Passed": counts["only_contract"],
            "Only Optimized Passed": counts["only_optimized"],
            "Contract Δ vs Vanilla": counts["contract_passed"] - counts["vanilla_passed"],
            "Optimized Δ vs Vanilla": counts["optimized_passed"] - counts["vanilla_passed"],
            "Optimized Δ vs Contract": counts["optimized_passed"] - counts["contract_passed"],
        })

    return rows


def last_error(trace: str) -> str:
    lines = [line.strip() for line in text(trace).splitlines() if line.strip()]
    markers = (
        "AssertionError",
        "TypeError",
        "ValueError",
        "KeyError",
        "AttributeError",
        "IndexError",
        "NameError",
        "SyntaxError",
        "ImportError",
        "Timeout",
    )
    return next((line for line in reversed(lines) if any(marker in line for marker in markers)), lines[-1] if lines else "")


def extract_expected_actual(trace: str, status: str = "") -> tuple[str, str]:
    trace = text(trace)

    for line in trace.splitlines():
        line = line.strip()
        patterns = [
            (r"AssertionError:\s*(.+?)\s*!=\s*(.+)$", "neq"),
            (r"AssertionError:\s*(.+?)\s+is not true", "true"),
            (r"AssertionError:\s*(.+?)\s+is not false", "false"),
            (r"AssertionError:\s*(.+?)\s+is not None", "none"),
            (r"AssertionError:\s*(.+?)\s+is not an instance of\s+(.+)$", "instance"),
            (r"AssertionError:\s*(.+?)\s+not found in\s+(.+)$", "contains"),
            (r"AssertionError:\s*(.+?)\s+unexpectedly found in\s+(.+)$", "not_contains"),
            (r"AssertionError:\s*(.+?)\s+!=\s+(.+?)\s+within\s+.+$", "almost"),
        ]

        for pattern, kind in patterns:
            match = re.search(pattern, line, re.I)
            if not match:
                continue

            if kind in {"neq", "almost"}:
                return short(match.group(2), 500), short(match.group(1), 500)
            if kind == "true":
                return "True", short(match.group(1), 500)
            if kind == "false":
                return "False", short(match.group(1), 500)
            if kind == "none":
                return "None", short(match.group(1), 500)
            if kind == "instance":
                return f"instance of {short(match.group(2), 250)}", short(match.group(1), 250)
            if kind == "contains":
                return f"contains {short(match.group(1), 250)}", short(match.group(2), 500)
            if kind == "not_contains":
                return f"does not contain {short(match.group(1), 250)}", short(match.group(2), 500)

    for prefix in ("Lists differ:", "Tuples differ:", "Dictionaries differ:", "Sets differ:"):
        match = re.search(re.escape(prefix) + r"\s*(.*?)\s*!=\s*(.*)", trace, re.DOTALL)
        if match:
            return short(match.group(2), 500), short(match.group(1), 500)

    error_line = last_error(trace)
    if status == "ERROR" or any(x in error_line for x in ("TypeError", "ValueError", "KeyError", "AttributeError", "IndexError")):
        return "No exception; successful execution", short(error_line, 500)

    return "", short(error_line, 500)


def failures(detail: JsonDict) -> list[JsonDict]:
    if detail.get("passed") is True:
        return []

    raw_type = text(detail.get("failure_type"))

    if raw_type in {"generation_failed", "missing_generation"}:
        return [{
            "status": "GENERATION_FAILED",
            "test_class": "",
            "test_name": "",
            "traceback": "",
            "error_line": text(detail.get("error")),
            "expected": "Valid complete Python class implementation",
            "actual": text(detail.get("error")),
        }]

    raw = text(as_dict(detail.get("evaluation")).get("stderr") or detail.get("error"))
    pattern = re.compile(
        r"=+\n(FAIL|ERROR):\s+(test[^\s]*)\s+\(([^)]+)\)\n-+\n(.*?)(?=\n=+\n|\n-+\nRan\s+\d+\s+tests?|\Z)",
        re.DOTALL,
    )

    rows: list[JsonDict] = []
    for match in pattern.finditer(raw):
        status = match.group(1)
        trace = match.group(4).strip()
        expected, actual = extract_expected_actual(trace, status)
        dotted = match.group(3)

        rows.append({
            "status": status,
            "test_class": dotted.split(".")[-2] if "." in dotted else dotted,
            "test_name": match.group(2),
            "traceback": short(trace, 900),
            "error_line": short(last_error(trace), 500),
            "expected": expected or "Expected behavior from benchmark test",
            "actual": actual or short(last_error(trace), 500),
        })

    if rows:
        return rows

    expected, actual = extract_expected_actual(raw)
    return [{
        "status": raw_type or "FAILED",
        "test_class": "",
        "test_name": "",
        "traceback": "",
        "error_line": short(detail.get("error") or raw, 700),
        "expected": expected or "Expected behavior from benchmark test",
        "actual": actual or short(detail.get("error") or raw, 700),
    }]


def test_metrics(detail: JsonDict, task: JsonDict) -> tuple[int | str, int | str]:
    metrics = metrics_of(detail)
    passed = metrics.get("passed")
    total = metrics.get("total")

    if passed is not None and total is not None:
        return int(passed), int(total)

    inferred_total = dataset_test_count(task)

    if detail.get("passed") is True:
        return inferred_total, inferred_total

    if inferred_total:
        return 0, inferred_total

    fail_count = len(failures(detail))
    return 0, fail_count if fail_count else ""


def norm(value: str) -> str:
    value = re.sub(r"^test_", "", text(value))
    value = re.sub(r"_\d+$", "", value)
    value = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", value)
    return re.sub(r"[^A-Za-z0-9_]+", "_", value).strip("_").lower()


def guess_method(test_name: str, method_names: list[str]) -> str:
    candidate = norm(test_name)
    lookup = {norm(method): method for method in method_names}

    if candidate in lookup:
        return lookup[candidate]

    return next((method for key, method in lookup.items() if candidate and (candidate in key or key in candidate)), candidate)


def generated_record(root: Path, detail: JsonDict | None) -> JsonDict:
    path = resolve(root, as_dict((detail or {}).get("generation")).get("path"))
    if not path or not path.exists():
        return {}
    data = load_json(path)
    return data if isinstance(data, dict) else {}


def code_from(record: JsonDict) -> str:
    return text(record.get("generated_code") or record.get("code"))


def contract_for(root: Path, run: JsonDict, task_id_value: str, folder: str, gen: JsonDict) -> JsonDict:
    if isinstance(gen.get("contract"), dict):
        return gen["contract"]

    provider = text(run["provider"])
    model = text(run["model"])

    candidates = []
    if folder == "rl_optimized_contract_guided":
        candidates.append(
            root
            / "outputs"
            / BENCHMARK
            / "rl_optimized_contracts"
            / safe_name(provider)
            / safe_name(model)
            / f"{safe_name(task_id_value)}.json"
        )

    candidates.append(
        root
        / "outputs"
        / BENCHMARK
        / "contracts"
        / safe_name(provider)
        / safe_name(model)
        / f"{safe_name(task_id_value)}_contract.json"
    )

    for path in candidates:
        if path.exists():
            return as_dict(load_json(path).get("contract"))

    return {}


def method_contracts(contract: JsonDict) -> dict[str, JsonDict]:
    return {
        text(row.get("method_name")): row
        for row in as_list(contract.get("method_contracts"))
        if isinstance(row, dict) and text(row.get("method_name"))
    }


def method_code(code: str, method: str) -> str:
    if not code or not method:
        return ""

    match = re.search(
        rf"(^[ \t]*(?:@\w+(?:\([^)]*\))?\s*\n[ \t]*)*def\s+{re.escape(method)}\s*\(.*?)(?=^[ \t]*(?:@\w|def\s+|class\s+)|\Z)",
        code,
        re.DOTALL | re.MULTILINE,
    )
    return short(match.group(1), 1200) if match else ""


def contract_type_from_evidence(
    detail: JsonDict,
    failure: JsonDict,
    task: JsonDict,
    method: str,
    code: str,
    contract: JsonDict,
) -> str:
    raw_type = text(detail.get("failure_type"))
    err = text(failure.get("error_line")).lower()
    expected = text(failure.get("expected")).lower()
    actual = text(failure.get("actual")).lower()
    trace = text(failure.get("traceback")).lower()
    task_text = " ".join([
        text(task.get("class_description")),
        text(task.get("class_constructor")),
        short(task.get("methods_info"), 3000),
    ]).lower()
    contract_text = short(contract, 3000).lower()
    method_text = method_code(code, method).lower()
    combined = " ".join([raw_type, err, expected, actual, trace, task_text, contract_text, method_text, text(method).lower()])

    if "changed method signatures" in combined or "staticmethod" in combined or "signature" in combined:
        return "interface"

    if raw_type in {"generation_failed", "missing_generation", "syntax_error", "import_error", "name_error", "timeout"}:
        return "not_contract_related"

    if "__init__" in combined or "constructor" in combined or "initializ" in combined:
        return "postcondition"

    if any(x in combined for x in ("invariant", "always", "must remain", "state valid")):
        return "invariant"

    if any(x in err for x in ("typeerror", "valueerror", "keyerror", "attributeerror", "indexerror", "zerodivisionerror")):
        if any(x in combined for x in ("invalid", "exception", "raise", "error input")):
            return "invalid_input_behavior"
        if any(x in combined for x in ("empty", "none", "null", "zero", "0", "missing", "boundary", "duplicate", "out of range")):
            return "edge_case"
        return "precondition"

    if any(x in combined for x in ("database", "sqlite", "cursor", "connection", "file", "path", "dependency", "calls", "interaction")):
        return "dependency_interaction"

    if any(x in combined for x in ("empty", "none", "null", "zero", "0", "missing", "duplicate", "boundary")):
        return "edge_case"

    if expected or actual:
        return "postcondition"

    return "not_contract_related"


def reason_and_solution(error_type: str, contract_type: str, failure: JsonDict) -> tuple[str, str]:
    if error_type == "Generation Error":
        return "The model did not produce a complete extractable implementation.", "Regenerate with strict class, signature, and executable-body constraints."
    if error_type == "Syntax Error":
        return "The generated code has syntax, import, or unresolved-name problems.", "Repair syntax, imports, and identifier references."
    if error_type == "Runtime Error":
        if contract_type == "precondition":
            return "The implementation rejects or crashes on an input/state case that should be handled.", "Relax wrong checks and implement the valid-input path."
        if contract_type == "edge_case":
            return "The implementation crashes on a boundary or unusual input/state case.", "Handle zero, empty, missing, duplicate, and None-like cases explicitly."
        return "The implementation raises an exception during benchmark execution.", "Guard the failing operation and implement the expected behavior."
    if error_type == "Timeout Error":
        return "The implementation likely contains inefficient or non-terminating logic.", "Replace unbounded loops with a bounded algorithm."
    if error_type == "Logical Error":
        if contract_type == "postcondition":
            return "The code executes but returns or stores a result different from the expected behavior.", "Correct the method logic to satisfy the expected postcondition."
        if contract_type == "invariant":
            return "The code executes but violates an expected persistent object property.", "Maintain the required invariant across constructor and method calls."
        if contract_type == "edge_case":
            return "The code executes but mishandles a boundary or special case.", "Add explicit handling for the failing edge case."
        if contract_type == "dependency_interaction":
            return "The code executes but mishandles a dependency or method interaction.", "Correct the interaction and preserve required side effects."
        return "The code executes but its behavior does not match the benchmark expectation.", "Align returned value, state effect, and formatting with the task."
    return text(failure.get("error_line")) or "The implementation does not satisfy benchmark behavior.", "Inspect the failed method and align it with task behavior."


def deterministic_analysis(
    detail: JsonDict,
    failure: JsonDict,
    task: JsonDict,
    method: str,
    code: str,
    contract: JsonDict,
) -> JsonDict:
    error_type = raw_to_error_type(detail.get("failure_type"), text(failure.get("error_line")))
    contract_type = contract_type_from_evidence(detail, failure, task, method, code, contract)
    reason, solve = reason_and_solution(error_type, contract_type, failure)
    return {
        "error_type": error_type,
        "contract_type": contract_type,
        "reason_for_failure": reason,
        "how_to_solve": solve,
    }


def analysis_prompt(
    model_label: str,
    pipeline: str,
    tid: str,
    class_name: str,
    failure: JsonDict,
    task: JsonDict,
    method: str,
    code: str,
    contract: JsonDict,
    base_analysis: JsonDict,
) -> str:
    payload = {
        "source_model": model_label,
        "pipeline": pipeline,
        "task_id": tid,
        "class_name": class_name,
        "failure": failure,
        "deterministic_error_type": base_analysis.get("error_type"),
        "deterministic_contract_type": base_analysis.get("contract_type"),
        "task_summary": short(task.get("class_description"), 1200),
        "failed_method_code": method_code(code, method),
        "related_contract": contract,
    }

    return f"""Return exactly one JSON object:
{{
  "error_type": "one of: {', '.join(ERROR_TYPES)}",
  "contract_type": "one of: {', '.join(CONTRACT_TYPES)}",
  "reason_for_failure": "max 35 words",
  "how_to_solve": "max 30 words"
}}

Rules:
- Never use raw evaluator labels such as wrong_answer.
- If raw failure_type is wrong_answer, classify it as Logical Error.
- Contract type must be one of the contract schema clause names only.

Use only this input:
{json.dumps(payload, ensure_ascii=False, indent=2)}
"""


def parse_object(raw: str) -> JsonDict:
    raw = text(raw)
    if not raw:
        return {}

    try:
        data = json.loads(raw)
        return data if isinstance(data, dict) else {}
    except json.JSONDecodeError:
        pass

    start = raw.find("{")
    end = raw.rfind("}")
    if start >= 0 and end > start:
        try:
            data = json.loads(raw[start:end + 1])
            return data if isinstance(data, dict) else {}
        except json.JSONDecodeError:
            return {}

    return {}


def normalize_analysis(raw: JsonDict, fallback: JsonDict) -> JsonDict:
    error_type = text(raw.get("error_type"))
    contract_type = text(raw.get("contract_type"))

    if error_type not in ERROR_TYPES:
        error_type = fallback["error_type"]

    if contract_type not in CONTRACT_TYPES:
        contract_type = fallback["contract_type"]

    return {
        "error_type": error_type,
        "contract_type": contract_type,
        "reason_for_failure": text(raw.get("reason_for_failure")) or fallback["reason_for_failure"],
        "how_to_solve": text(raw.get("how_to_solve")) or fallback["how_to_solve"],
    }


def analyze_llm(client: Any, prompt: str, temperature: float, max_tokens: int) -> JsonDict:
    raw, _ = call_chat_model(
        client=client,
        model=ANALYZER_MODEL,
        messages=[
            {"role": "system", "content": "You are an expert Python benchmark failure analyst. Return one valid JSON object only."},
            {"role": "user", "content": prompt},
        ],
        temperature=temperature,
        max_tokens=max_tokens,
        json_mode=True,
    )

    parsed = parse_object(raw)
    if parsed:
        return parsed

    raise ValueError(short(raw, 500))


def cache_load(path: Path) -> dict[str, JsonDict]:
    if not path.exists():
        return {}

    out: dict[str, JsonDict] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        try:
            row = json.loads(line)
        except json.JSONDecodeError:
            continue
        if text(row.get("_key")):
            out[text(row["_key"])] = row
    return out


def cache_append(path: Path, key: str, row: JsonDict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as file:
        file.write(json.dumps({"_key": key, **row}, ensure_ascii=False) + "\n")


def detail_rows_for(
    root: Path,
    run: JsonDict,
    model_label: str,
    loaded: dict[str, tuple[str, Path | None, list[JsonDict]]],
    dataset: dict[str, JsonDict],
    use_llm: bool,
    client: Any,
    cache: dict[str, JsonDict],
    cache_path: Path,
    max_llm_rows: int | None,
    temperature: float,
    max_tokens: int,
    delay: float,
) -> list[JsonDict]:
    rows: list[JsonDict] = []
    llm_used = 0

    for pipeline_label, (folder, _path, details) in loaded.items():
        for detail in details:
            if detail.get("passed") is True:
                continue

            tid = text(detail.get("task_id"))
            task = dataset.get(tid, {})
            method_names = methods(task)
            class_label = text(detail.get("class_name") or task.get("class_name") or task.get("entry_point"))
            gen = generated_record(root, detail)
            code = code_from(gen)
            contract = method_contracts(contract_for(root, run, tid, folder, gen))
            tests_passed, tests_total = test_metrics(detail, task)

            for failure in failures(detail):
                method = guess_method(text(failure.get("test_name")), method_names)
                base_analysis = deterministic_analysis(detail, failure, task, method, code, contract.get(method, {}))
                analysis = base_analysis

                key = "|".join([
                    model_label,
                    pipeline_label,
                    tid,
                    text(failure.get("test_class")),
                    text(failure.get("test_name")),
                ])

                if use_llm and (max_llm_rows is None or llm_used < max_llm_rows):
                    if key in cache:
                        analysis = normalize_analysis(cache[key], base_analysis)
                    else:
                        prompt = analysis_prompt(
                            model_label=model_label,
                            pipeline=pipeline_label,
                            tid=tid,
                            class_name=class_label,
                            failure=failure,
                            task=task,
                            method=method,
                            code=code,
                            contract=contract.get(method, {}),
                            base_analysis=base_analysis,
                        )
                        try:
                            analysis = normalize_analysis(analyze_llm(client, prompt, temperature, max_tokens), base_analysis)
                        except Exception as exc:
                            analysis = base_analysis
                            analysis["reason_for_failure"] = (
                                f"{analysis.get('reason_for_failure', '')} "
                                f"LLM analyzer failed: {short(exc, 160)}"
                            ).strip()
                        cache_append(cache_path, key, analysis)
                        llm_used += 1
                        if delay:
                            time.sleep(delay)

                rows.append({
                    "Source Model": model_label,
                    "Pipeline": pipeline_label,
                    "Task ID": tid,
                    "Class Name": class_label,
                    "Methods List": ", ".join(method_names),
                    "Passed": False,
                    "Tests Passed": tests_passed,
                    "Tests Total": tests_total,
                    "Failed Method / Test": f"{method} / {failure.get('test_class')}.{failure.get('test_name')}".strip(" /."),
                    "Expected Output": text(failure.get("expected")),
                    "Actual Output": text(failure.get("actual")),
                    "Error Type": analysis["error_type"],
                    "Raw Failure Type": text(detail.get("failure_type")),
                    "Contract Type": analysis["contract_type"],
                    "Reason for Failure": analysis["reason_for_failure"],
                    "How to Solve": analysis["how_to_solve"],
                })

    return rows


def load_all(root: Path) -> dict[str, dict[str, tuple[str, Path | None, list[JsonDict]]]]:
    all_runs: dict[str, dict[str, tuple[str, Path | None, list[JsonDict]]]] = {}

    for run in SOURCE_RUNS:
        label = text(run["label"])
        all_runs[label] = {}

        for pipeline_label, folder in PIPELINES:
            required = pipeline_label != "Feedback-optimized"
            path = find_details(root, folder, run, required=required)

            if not path:
                all_runs[label][pipeline_label] = (folder, None, [])
                continue

            data = load_json(path)
            if not isinstance(data, list):
                raise ValueError(f"Details file must be a JSON list: {path}")

            all_runs[label][pipeline_label] = (folder, path, data)

    return all_runs


def write_summary_sheet(writer: pd.ExcelWriter, rows: dict[str, list[JsonDict]]) -> None:
    sheet_name = "Summary"
    start = 0

    for title, data in rows.items():
        pd.DataFrame([[title]]).to_excel(writer, sheet_name=sheet_name, index=False, header=False, startrow=start)
        start += 1

        df = pd.DataFrame(data)
        df.to_excel(writer, sheet_name=sheet_name, index=False, startrow=start)
        start += len(df) + 3


def style_workbook(writer: pd.ExcelWriter) -> None:
    title_fill = PatternFill("solid", fgColor="D9EAF7")
    header_fill = PatternFill("solid", fgColor="E2F0D9")

    for _sheet_name, ws in writer.sheets.items():
        ws.freeze_panes = "A2"

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        for row in ws.iter_rows():
            values = [cell.value for cell in row if cell.value not in (None, "")]
            if len(values) == 1 and isinstance(values[0], str) and values[0] in {
                "Pipeline Performance",
                "Pairwise Comparisons",
                "Three-way Gains",
            }:
                for cell in row:
                    cell.font = Font(bold=True)
                    cell.fill = title_fill

        for row in ws.iter_rows():
            first = text(row[0].value) if row else ""
            if first in {"Section", "Source Model"}:
                for cell in row:
                    cell.font = Font(bold=True)
                    cell.fill = header_fill

        for col in ws.columns:
            letter = col[0].column_letter
            width = max(12, max(min(len(str(cell.value or "")), 70) for cell in col[:250]) + 2)
            ws.column_dimensions[letter].width = width


def write_excel(path: Path, summary_sections: dict[str, list[JsonDict]], detail_rows: list[JsonDict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        write_summary_sheet(writer, summary_sections)
        pd.DataFrame(detail_rows, columns=DETAIL_COLUMNS).to_excel(writer, sheet_name="Failure Detail", index=False)
        style_workbook(writer)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate combined ClassEval report for both configured source models.")
    parser.add_argument("--root", default=".")
    parser.add_argument("--dataset-file")
    parser.add_argument("--out", default="results/classeval/analysis/classeval_two_model_report.xlsx")
    parser.add_argument("--use-llm", action="store_true")
    parser.add_argument("--max-llm-rows", type=int, default=0)
    parser.add_argument("--temperature", type=float, default=0.0)
    parser.add_argument("--max-tokens", type=int, default=500)
    parser.add_argument("--delay", type=float, default=0.0)
    parser.add_argument("--cache")
    parser.add_argument("--overwrite-cache", action="store_true")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    root = Path(args.root).resolve()
    dataset = load_dataset(root, args.dataset_file)
    all_runs = load_all(root)

    use_llm = bool(args.use_llm and args.max_llm_rows != 0)
    client = get_client(ANALYZER_PROVIDER) if use_llm else None

    cache_path = (
        Path(args.cache)
        if args.cache
        else root / "results" / BENCHMARK / "analysis" / f"failure_analysis_cache_{safe_name(ANALYZER_PROVIDER)}_{safe_name(ANALYZER_MODEL)}.jsonl"
    )
    cache = {} if args.overwrite_cache else cache_load(cache_path)

    detail_rows: list[JsonDict] = []

    for run in SOURCE_RUNS:
        label = text(run["label"])
        detail_rows.extend(
            detail_rows_for(
                root=root,
                run=run,
                model_label=label,
                loaded=all_runs[label],
                dataset=dataset,
                use_llm=use_llm,
                client=client,
                cache=cache,
                cache_path=cache_path,
                max_llm_rows=args.max_llm_rows if args.max_llm_rows > 0 else None,
                temperature=args.temperature,
                max_tokens=args.max_tokens,
                delay=args.delay,
            )
        )

    summary_sections = {
        "Pipeline Performance": pipeline_summary_rows(all_runs, detail_rows),
        "Pairwise Comparisons": comparison_rows(all_runs),
        "Three-way Gains": three_way_rows(all_runs),
    }

    write_excel(Path(args.out), summary_sections, detail_rows)

    print(f"Output Excel: {args.out}")
    print(f"Analyzer: {ANALYZER_PROVIDER} / {ANALYZER_MODEL}" if use_llm else "Analyzer: disabled; deterministic report used")

    for label, loaded in all_runs.items():
        print(f"\n{label}")
        for pipeline_label, (_folder, path, rows) in loaded.items():
            print(f"  {pipeline_label}: {len(rows)} rows from {path}")


if __name__ == "__main__":
    main()